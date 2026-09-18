"""Align the board-related parts of act101/102/103 with the shared กิจกรรมที่ 7 table.

Reference: data_dictionary_board_submission V2 (tbl_board_submission, shared by every activity).
Scope (agreed): board-related tables/columns only; everything else keeps its current names.
  * data the board sends back now lives in tbl_board_submission -> duplicated LAW columns/tables removed
  * one tbl_law{n}_board_link table per sub-activity records each submission (tbs_id)
  * board tables that stay in LAW are renamed to tbl_law{n}_* with a 3-letter column prefix
    and the standard audit columns (created_datetime / created_by INTEGER / ...)
  * tbl_board_submission is added to each diagram as an external reference (no DDL, no FK constraint)

Usage: python apply_board_standard.py "<path to data_dictionary_board_submission V2.xlsx>"
Always starts from backup-before-board-standard/, so it is safe to re-run; then run build_db_docs.py.
"""
import copy
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).parent
BOARD_XLSX = Path(sys.argv[1])
EXT_GROUP = {"code": "EXT7", "label": "กิจกรรมที่ 7 (ตารางกลาง — อ้างอิงภายนอก)"}
LINK_GROUP = {"code": "BOARD", "label": "ส่ง/รับ กิจกรรมที่ 7 (มาตรฐานกลาง)"}
LEGACY_AUDIT = {"created_at", "created_by", "updated_at", "updated_by"}


def col(name, ctype, label_th, description, pk=False, fk=None, nullable=True, default=None, source="system"):
    return {"name": name, "type": ctype, "pk": pk, "fk": fk, "nullable": nullable, "default": default,
            "unique": False, "label_th": label_th, "source": source, "description": description}


def audit_cols():
    return [
        col("created_datetime", "timestamptz", "วันเวลาที่สร้างข้อมูล", "Standard audit column (กิจกรรมที่ 7 convention)",
            nullable=False, default="now()"),
        col("created_by", "integer", "ผู้สร้างข้อมูล", "Staff id from the user system"),
        col("updated_datetime", "timestamptz", "วันเวลาแก้ไขล่าสุด", "Standard audit column"),
        col("updated_by", "integer", "ผู้แก้ไขล่าสุด", "Staff id from the user system"),
    ]


def board_submission_table():
    """tbl_board_submission read from the กิจกรรมที่ 7 data dictionary (external reference only)."""
    ws = load_workbook(BOARD_XLSX, data_only=True).worksheets[0]
    columns = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not isinstance(row[0], int):
            continue
        _, group, name, ctype, size, key, null, desc = row[:8]
        ctype = ctype.lower()
        if ctype in ("varchar", "char") and size not in (None, "-"):
            ctype = f"{ctype}({size})"
        elif ctype == "varchar[]":
            ctype = f"varchar({size})[]"
        columns.append(col(name, ctype, group, desc, pk=(key == "PK"), nullable=(null != "ไม่ได้"),
                           source="data_dictionary_board_submission V2"))
    return {"name": "tbl_board_submission", "label_th": "เรื่องที่ส่งเข้าคณะกรรมการ ป.ป.ท. และผลมติ (ตารางกลาง)",
            "group": EXT_GROUP["code"], "external": True,
            "description": "Shared กิจกรรมที่ 7 table (not created by LAW). LAW links to it by tbs_id without a foreign key.",
            "columns": columns, "unique_constraints": [], "check_constraints": [], "indexes": []}


def link_table(n, send_points, extra_cols):
    p = f"lb{n}_"
    points = ", ".join(f"'{s}'" for s in send_points)
    return {
        "name": f"tbl_law{n}_board_link",
        "label_th": "การส่งเรื่องเข้ากิจกรรมที่ 7 (เชื่อม tbl_board_submission)",
        "group": LINK_GROUP["code"],
        "description": ("One row per submission from this sub-activity to กิจกรรมที่ 7; "
                        "the resolution itself is read from tbl_board_submission."),
        "columns": [
            col(f"{p}id", "bigserial", "รหัส", "Primary key", pk=True, nullable=False),
            col(f"{p}case_id", "bigint", "แฟ้มคดี", "Case sent to the board", fk="case_file.id", nullable=False),
            *extra_cols(p),
            col(f"{p}send_point", "varchar(10)", "จุดส่ง", f"Handoff point id ({points})", nullable=False,
                source="ecmis-activity10-act7-handoffs.xlsx"),
            col(f"{p}round_no", "smallint", "ครั้งที่ส่ง", "Resubmission counter for the same point",
                nullable=False, default="1"),
            col(f"{p}send_step_code", "varchar(10)", "ขั้นตอนส่ง (LAW)", "TO-BE LAW step that submits",
                nullable=False, source="TO-BE UserFlow V1.0.1"),
            col(f"{p}receive_step_code", "varchar(10)", "ขั้นตอนรับผล (LAW)", "TO-BE LAW step that receives the resolution",
                source="TO-BE UserFlow V1.0.1"),
            col(f"{p}tbs_id", "bigint", "รหัสเรื่องเสนอ (กิจกรรมที่ 7)",
                "tbl_board_submission.tbs_id (cross-module, no FK constraint)",
                fk="tbl_board_submission.tbs_id", nullable=False),
            col(f"{p}submission_no", "varchar(50)", "เลขที่เรื่องเสนอ", "Copy of tbs_submission_no for display"),
            col(f"{p}matter_type", "varchar(3)", "ประเภทเรื่องที่เสนอ",
                "tbs_matter_type code (LAW block still to be registered with กิจกรรมที่ 7)", nullable=False),
            *audit_cols(),
        ],
        "unique_constraints": [[f"{p}tbs_id"], [f"{p}case_id", f"{p}send_point", f"{p}round_no"]],
        "check_constraints": [f"{p}send_point in ({points})", f"{p}round_no >= 1"],
        "indexes": [[f"{p}case_id"]],
    }


def table(schema, name):
    return next(t for t in schema["tables"] if t["name"] == name)


def rename_table(t, new_name, prefix, label_th, drop=()):
    """Rename a LAW board table to the standard: tbl_ name, column prefix, standard audit columns."""
    t["name"], t["label_th"] = new_name, label_th
    kept = [c for c in t["columns"] if c["name"] not in drop and c["name"] not in LEGACY_AUDIT]
    renamed = {}
    for c in kept:
        renamed[c["name"]] = prefix + c["name"]
        c["name"] = prefix + c["name"]
    t["columns"] = kept + audit_cols()
    for key in ("unique_constraints", "indexes"):
        t[key] = [g for g in ([renamed.get(x, x) for x in grp if x not in drop] for grp in t.get(key, [])) if g]
    t["check_constraints"] = [ck for ck in t.get("check_constraints", []) if not any(d in ck for d in drop)]
    t["group"] = LINK_GROUP["code"]
    return renamed


def drop_tables(schema, names):
    schema["tables"] = [t for t in schema["tables"] if t["name"] not in names]
    schema["enums"] = [e for e in schema["enums"] if e["table"] not in names]


def drop_columns(schema, table_name, names):
    t = table(schema, table_name)
    t["columns"] = [c for c in t["columns"] if c["name"] not in names]
    for key in ("unique_constraints", "indexes"):
        t[key] = [g for g in ([x for x in grp if x not in names] for grp in t.get(key, [])) if g]


def retarget_fks(schema, old_table, new_table, renamed):
    for t in schema["tables"]:
        for c in t["columns"]:
            if c.get("fk") and c["fk"].split(".")[0] == old_table:
                old_col = c["fk"].split(".")[1]
                c["fk"] = f"{new_table}.{renamed.get(old_col, old_col)}"


def replace_in_flow(schema, mapping):
    for f in schema["flow_map"]:
        for key in ("reads", "writes"):
            out = []
            for x in f.get(key, []):
                for y in mapping.get(x, [x]):
                    if y not in out:
                        out.append(y)
            f[key] = out


def add_board_tables(schema, tables):
    codes = {g["code"] for g in schema["groups"]}
    schema["groups"] += [g for g in (LINK_GROUP, EXT_GROUP) if g["code"] not in codes]
    schema["tables"].extend(tables)


def board_targets(n):
    return [f"tbl_law{n}_board_link", "tbl_board_submission"]


def apply_101(s, ext):
    drop_tables(s, {"board_resolution", "executive_sign"})
    link = link_table(1, ["S101-1", "S101-2"], lambda p: [
        col(f"{p}dispatch_id", "bigint", "การออกเลขส่ง", "admin_dispatch row that sent this submission (pages 09 / 16)",
            fk="admin_dispatch.id"),
    ])
    add_board_tables(s, [link, ext])
    replace_in_flow(s, {"board_resolution": board_targets(1), "executive_sign": board_targets(1)})
    s["notes"].append(
        "[มาตรฐานกิจกรรมที่ 7] ลบตาราง board_resolution และ executive_sign — มติ/ผลลงนามรับกลับจาก tbl_board_submission "
        "(tbs_resolution_result/ref/date/detail/doc_ref/received_datetime) ผ่าน tbl_law1_board_link "
        "(S101-1 LAW0010, S101-2 LAW0019). ชื่อผู้บริหารที่ลงนามไม่มีฟิลด์ในตารางกลาง — ถือเป็นข้อมูลภายในกิจกรรมที่ 7")


def apply_102(s, ext):
    drop_columns(s, "case_file", {"l2_board_approval_ref", "l2_board_approval_date"})
    ack = table(s, "board_round2")
    renamed = rename_table(ack, "tbl_law2_board_ack", "ba2_", "การรับทราบมติคณะกรรมการ ป.ป.ท. (รอบ 2)",
                           drop={"approval_ref", "approval_date", "resolution_text"})
    ack["description"] = ("LAW-internal acknowledgement of the round-2 board resolution; "
                          "the resolution itself is in tbl_board_submission.")
    retarget_fks(s, "board_round2", "tbl_law2_board_ack", renamed)
    link = link_table(2, ["S102-1", "S102-2", "S102-3"], lambda p: [
        col(f"{p}appeal_id", "bigint", "คำอุทธรณ์", "Set for appeal submissions (S102-3)", fk="appeal.id"),
    ])
    add_board_tables(s, [link, ext])
    replace_in_flow(s, {"board_round2": ["tbl_law2_board_ack", *board_targets(2)]})
    for e in s["enums"]:
        if e["table"] in ("resolution_type", "appeal_board_resolution_type"):
            block = "L1" if e["table"] == "resolution_type" else "L2"
            for i, row in enumerate(e["rows"], start=1):
                row.setdefault("extra", {})["tbs_resolution_result"] = "901" if row["code"] == "OTHER" else f"{block}{i}"
    s["notes"] += [
        "[มาตรฐานกิจกรรมที่ 7] ลบ case_file.l2_board_approval_ref/date และ board_round2.approval_ref/approval_date/resolution_text — "
        "อ่านจาก tbl_board_submission.tbs_resolution_ref/date/detail; board_round2 เปลี่ยนชื่อเป็น tbl_law2_board_ack "
        "(คงเฉพาะการรับทราบภายใน)",
        "[มาตรฐานกิจกรรมที่ 7] ตาราง resolution คงไว้สำหรับมติคณะอนุกรรมการกลั่นกรอง (ภายใน) และ case_state ซึ่งตารางกลางไม่มีฟิลด์รองรับ; "
        "มติรอบเลขาธิการ/รองเลขาธิการ (S102-1) และมติอุทธรณ์ (S102-3) รับจาก tbl_board_submission "
        "— appeal_step แถว appeal-10/11 ไม่ต้องเก็บเลขมติ/ผลซ้ำ",
        "[มาตรฐานกิจกรรมที่ 7] เสนอรหัส tbs_resolution_result ของ LAW (ยังไม่ได้ลงทะเบียน): resolution_type → L11 เปิดเผย, "
        "L12 บางส่วน, L13 ไม่อนุญาต, OTHER → 901; appeal_board_resolution_type → L21/L22/L23 (ดูคอลัมน์ Extra ในชีต Enums)",
    ]


def apply_103(s, ext):
    drop_tables(s, {"board_resolution", "board_propose_result"})
    notice = table(s, "resolution_notice")
    renamed = rename_table(notice, "tbl_law3_board_notice", "bn3_", "การรับแจ้งและกระจายผลมติคณะกรรมการ ป.ป.ท.",
                           drop={"doc_no", "doc_date"})
    retarget_fks(s, "resolution_notice", "tbl_law3_board_notice", renamed)
    recipient = table(s, "resolution_notice_recipient")
    renamed_r = rename_table(recipient, "tbl_law3_board_notice_recipient", "br3_", "ผู้รับแจ้งผลมติ")
    retarget_fks(s, "resolution_notice_recipient", "tbl_law3_board_notice_recipient", renamed_r)
    add_board_tables(s, [link_table(3, ["S103-1", "S103-2"], lambda p: []), ext])
    replace_in_flow(s, {
        "board_resolution": board_targets(3),
        "board_propose_result": board_targets(3),
        "resolution_notice": ["tbl_law3_board_notice", *board_targets(3)],
        "resolution_notice_recipient": ["tbl_law3_board_notice_recipient"],
    })
    s["notes"].append(
        "[มาตรฐานกิจกรรมที่ 7] ลบตาราง board_resolution และ board_propose_result, ลบ resolution_notice.doc_no/doc_date — "
        "อ่านจาก tbl_board_submission; resolution_notice(+recipient) เปลี่ยนชื่อเป็น tbl_law3_board_notice(+_recipient). "
        "R103-2A/2B แยกตาม tbs_resolution_result (เสนอ L31 อุทธรณ์ / L32 ไม่อุทธรณ์)")


def main():
    ext = board_submission_table()
    for n, apply in ((1, apply_101), (2, apply_102), (3, apply_103)):
        name = f"schema-10-{n}.json"
        schema = json.loads((HERE / "backup-before-board-standard" / name).read_text(encoding="utf-8"))
        apply(schema, copy.deepcopy(ext))
        (HERE / name).write_text(json.dumps(schema, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"{schema['schema']}: {len(schema['tables'])} tables")


if __name__ == "__main__":
    main()
