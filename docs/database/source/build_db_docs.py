"""Build the E-CMIS activity 10 DB deliverables from schema-10-*.json.

Usage: python build_db_docs.py OUT_DIR [schema.json ...]
Outputs (in OUT_DIR):
  ecmis-activity10-er.drawio             one diagram page per activity
  ecmis-activity10-data-dictionary.xlsx
  ecmis-activity10-schema.sql            CREATE statements generated from the same JSON
"""
import json
import sys
from pathlib import Path
from xml.sax.saxutils import escape as _escape


def escape(text):
    """Escape for use inside a double-quoted XML attribute."""
    return _escape(text, {'"': "&quot;"})

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SCRATCH = Path(__file__).parent
OUT_DIR = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"D:\Mercil\E-CMIS-A4\docs\database")
INPUTS = [Path(p) for p in sys.argv[2:]] or [SCRATCH / f"schema-10-{n}.json" for n in (1, 2, 3)]

FONT = "Arial"
ROW_H = 22
HEADER_H = 52
CHAR_W = 7
MIN_W = 220
KEY_W = 34
COL_GAP = 90
ROW_GAP = 40
TOP = 130
MAX_COL_H = 2600

GROUP_COLORS = ["#dae8fc", "#d5e8d4", "#fff2cc", "#f8cecc", "#e1d5e7", "#ffe6cc", "#f5f5f5", "#d0e0e3"]
GROUP_STROKES = ["#6c8ebf", "#82b366", "#d6b656", "#b85450", "#9673a6", "#d79b00", "#666666", "#4a7c8c"]


def load(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


# ---------------------------------------------------------------- drawio
def row_text(col):
    key = "PK" if col.get("pk") else ("FK" if col.get("fk") else "")
    not_null = "" if col.get("nullable", True) or col.get("pk") else " NN"
    return key, f"{col['name']} : {col['type']}{not_null}"


def table_size(t):
    texts = [row_text(c)[1] for c in t["columns"]]
    longest = max([len(x) for x in texts] + [len(t["name"]) + 4])
    width = max(MIN_W, longest * CHAR_W + KEY_W + 36, len(t.get("label_th", "")) * 6 + 40)
    return width, HEADER_H + ROW_H * len(t["columns"])


def layout(schema):
    """One column per group (in `groups` order); a group wraps to a new column when too tall."""
    group_order = [g["code"] for g in schema.get("groups", [])]
    by_group = {}
    for t in schema["tables"]:
        by_group.setdefault(t.get("group", "OTHER"), []).append(t)
    group_order += [g for g in by_group if g not in group_order]

    positions, group_x, x = {}, {}, 40
    for g in group_order:
        tables = by_group.get(g, [])
        if not tables:
            continue
        start_x = x
        col_w, y = 0, TOP
        for t in tables:
            w, h = table_size(t)
            if y > TOP and y + h > MAX_COL_H:
                x += col_w + COL_GAP
                col_w, y = 0, TOP
            positions[t["name"]] = (x, y, w, h)
            col_w = max(col_w, w)
            y += h + ROW_GAP
        group_x[g] = (start_x, x + col_w - start_x)  # (left, width) so the label never spills into the next group
        x += col_w + COL_GAP
    return positions, group_order, group_x


def assert_no_overlap(positions):
    boxes = list(positions.items())
    for i, (a, (ax, ay, aw, ah)) in enumerate(boxes):
        for b, (bx, by, bw, bh) in boxes[i + 1:]:
            if ax < bx + bw and bx < ax + aw and ay < by + bh and by < ay + ah:
                raise ValueError(f"layout overlap: {a} / {b}")


def diagram_page(schema, page_idx):
    positions, group_order, group_x = layout(schema)
    assert_no_overlap(positions)
    color_of = {g: i % len(GROUP_COLORS) for i, g in enumerate(group_order)}
    group_label = {g["code"]: g["label"] for g in schema.get("groups", [])}
    p = f"p{page_idx}"
    cells = [f'<mxCell id="{p}-0"/>', f'<mxCell id="{p}-1" parent="{p}-0"/>']

    title = escape(f"{schema['activity']} {schema.get('title_th', '')}  —  schema {schema['schema']}")
    cells.append(
        f'<mxCell id="{p}-title" value="{title}" style="text;html=1;fontSize=22;fontStyle=1;fontFamily={FONT};align=left;verticalAlign=middle;" vertex="1" parent="{p}-1">'
        f'<mxGeometry x="40" y="20" width="1200" height="40" as="geometry"/></mxCell>'
    )

    for g, (gx, gw) in group_x.items():
        ci = color_of[g]
        label = escape(group_label.get(g, g))
        cells.append(
            f'<mxCell id="{p}-g-{g}" value="{label}" style="text;html=1;whiteSpace=wrap;fontSize=13;fontStyle=1;fontFamily={FONT};align=left;verticalAlign=bottom;fontColor={GROUP_STROKES[ci]};" vertex="1" parent="{p}-1">'
            f'<mxGeometry x="{gx}" y="{TOP - 56}" width="{gw}" height="48" as="geometry"/></mxCell>'
        )

    row_ids, pk_row = {}, {}
    for t in schema["tables"]:
        x, y, w, h = positions[t["name"]]
        ci = color_of[t.get("group", "OTHER")]
        tid = f"{p}-t-{t['name']}"
        sub = t.get("label_th", "")
        header = escape(f'{t["name"]}<br><font style="font-size:10px">{_escape(sub)}</font>')
        cells.append(
            f'<mxCell id="{tid}" value="{header}" style="shape=table;startSize={HEADER_H};container=1;collapsible=0;childLayout=tableLayout;fixedRows=1;rowLines=0;fontStyle=1;align=center;resizeLast=1;html=1;whiteSpace=wrap;fontFamily={FONT};fontSize=13;fillColor={GROUP_COLORS[ci]};strokeColor={GROUP_STROKES[ci]};swimlaneFillColor=#ffffff;" vertex="1" parent="{p}-1">'
            f'<mxGeometry x="{x}" y="{y}" width="{w}" height="{h}" as="geometry"/></mxCell>'
        )
        for r, col in enumerate(t["columns"]):
            key, text = row_text(col)
            rid = f"{tid}-r{r}"
            row_ids[(t["name"], col["name"])] = rid
            if col.get("pk"):
                pk_row.setdefault(t["name"], rid)
            bold = 5 if col.get("pk") else 0
            cells.append(
                f'<mxCell id="{rid}" value="" style="shape=tableRow;horizontal=0;startSize=0;swimlaneHead=0;swimlaneBody=0;fillColor=none;collapsible=0;dropTarget=0;points=[[0,0.5],[1,0.5]];portConstraint=eastwest;top=0;left=0;right=0;bottom={1 if col.get("pk") else 0};fontFamily={FONT};" vertex="1" parent="{tid}">'
                f'<mxGeometry y="{HEADER_H + r * ROW_H}" width="{w}" height="{ROW_H}" as="geometry"/></mxCell>'
            )
            cells.append(
                f'<mxCell id="{rid}-k" value="{key}" style="shape=partialRectangle;connectable=0;fillColor=none;top=0;left=0;bottom=0;right=0;fontStyle={bold};overflow=hidden;whiteSpace=wrap;html=1;fontFamily={FONT};fontSize=11;fontColor=#555555;" vertex="1" parent="{rid}">'
                f'<mxGeometry width="{KEY_W}" height="{ROW_H}" as="geometry"><mxRectangle width="{KEY_W}" height="{ROW_H}" as="alternateBounds"/></mxGeometry></mxCell>'
            )
            cells.append(
                f'<mxCell id="{rid}-v" value="{escape(text)}" style="shape=partialRectangle;connectable=0;fillColor=none;top=0;left=0;bottom=0;right=0;align=left;spacingLeft=6;fontStyle={bold};overflow=hidden;whiteSpace=wrap;html=1;fontFamily={FONT};fontSize=11;" vertex="1" parent="{rid}">'
                f'<mxGeometry x="{KEY_W}" width="{w - KEY_W}" height="{ROW_H}" as="geometry"><mxRectangle width="{w - KEY_W}" height="{ROW_H}" as="alternateBounds"/></mxGeometry></mxCell>'
            )

    edge_n = 0
    for t in schema["tables"]:
        for col in t["columns"]:
            fk = col.get("fk")
            if not fk:
                continue
            ref_table, _, ref_col = fk.partition(".")
            target = row_ids.get((ref_table, ref_col)) or pk_row.get(ref_table)
            if not target:
                print(f"  WARN {schema['schema']}.{t['name']}.{col['name']} -> {fk}: target not found")
                continue
            end = "ERzeroToOne" if col.get("nullable", True) else "ERmandOne"
            edge_n += 1
            cells.append(
                f'<mxCell id="{p}-e{edge_n}" value="" style="edgeStyle=entityRelationEdgeStyle;fontSize=12;html=1;endArrow={end};startArrow=ERmany;rounded=0;strokeColor=#7f8c8d;opacity=70;" edge="1" parent="{p}-1" source="{row_ids[(t["name"], col["name"])]}" target="{target}">'
                f'<mxGeometry relative="1" as="geometry"/></mxCell>'
            )

    name = escape(f"{schema['activity']} {schema['schema']}")
    xml = (f'<diagram id="act-{page_idx}" name="{name}"><mxGraphModel dx="1600" dy="900" grid="1" gridSize="10" '
           f'guides="1" tooltips="1" connect="1" arrows="1" fold="1" page="0" pageScale="1" pageWidth="1169" '
           f'pageHeight="827" math="0" shadow="0"><root>{"".join(cells)}</root></mxGraphModel></diagram>')
    return xml, edge_n


def build_drawio(schemas, path):
    pages, total_edges = [], 0
    for i, s in enumerate(schemas, start=1):
        xml, n = diagram_page(s, i)
        pages.append(xml)
        total_edges += n
    path.write_text(f'<mxfile host="app.diagrams.net">{"".join(pages)}</mxfile>', encoding="utf-8")
    return total_edges


# ---------------------------------------------------------------- xlsx
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10)
PK_FILL = PatternFill("solid", fgColor="FFF2CC")
FK_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_sheet(ws, headers, rows, widths, fills=None):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, row in enumerate(rows):
        ws.append(row)
        fill = fills[i] if fills else None
        for c in ws[ws.max_row]:
            c.font, c.border = BODY_FONT, BORDER
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if fill:
                c.fill = fill
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_readme(ws):
    ws.title = "README"
    rows = [
        ["E-CMIS กิจกรรมที่ 10 — Data dictionary (PostgreSQL)"],
        [],
        ["What this is", "Database design that persists every value passing through the 10.1 / 10.2 / 10.3 workflows, derived page by page from the activity10 mockup."],
        ["Structure", "Each activity is a fully separate PostgreSQL schema (act101, act102, act103) with no cross-schema foreign keys."],
        ["Sheets", "Tables = one row per table · 10.x Columns = every column · Enums = lookup rows · Flow map = tables each page reads/writes · Notes = design decisions and gaps found in the mockup."],
        ["Row colours", "Yellow = primary key column · Blue = foreign key column."],
        ["Workflow history", "case_transition records every submit / send-back / branch, so review rounds are never overwritten."],
        ["Companion files", "ecmis-activity10-er.drawio (one page per activity) and ecmis-activity10-schema.sql, generated from the same source."],
        ["Editing", "Generated from schema-10-*.json — change the JSON and regenerate rather than editing this workbook by hand."],
    ]
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(name=FONT, size=16, bold=True)
    for row in ws.iter_rows(min_row=3):
        row[0].font = Font(name=FONT, bold=True)
        for c in row[1:]:
            c.font = BODY_FONT
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 110


def build_xlsx(schemas, path):
    wb = Workbook()
    build_readme(wb.active)

    t_rows = []
    for s in schemas:
        glabel = {g["code"]: g["label"] for g in s.get("groups", [])}
        for t in s["tables"]:
            refs = sorted({c["fk"].split(".")[0] for c in t["columns"] if c.get("fk")})
            t_rows.append([s["activity"], s["schema"], t["name"], t.get("label_th", ""),
                           glabel.get(t.get("group"), t.get("group", "")), t.get("description", ""),
                           len(t["columns"]), ", ".join(refs)])
    write_sheet(wb.create_sheet("Tables"),
                ["Activity", "Schema", "Table", "Thai name", "Group", "Description", "Columns", "References"],
                t_rows, [9, 10, 30, 30, 22, 60, 9, 40])

    for s in schemas:
        rows, fills = [], []
        for t in s["tables"]:
            for i, c in enumerate(t["columns"], start=1):
                rows.append([s["schema"], t["name"], i, c["name"], c["type"],
                             "Y" if c.get("pk") else "", c.get("fk") or "",
                             "Y" if c.get("nullable", True) else "N",
                             "" if c.get("default") is None else str(c["default"]),
                             "Y" if c.get("unique") else "", c.get("label_th", ""),
                             c.get("source", ""), c.get("description", "")])
                fills.append(PK_FILL if c.get("pk") else (FK_FILL if c.get("fk") else None))
        write_sheet(wb.create_sheet(f"{s['activity']} Columns"),
                    ["Schema", "Table", "#", "Column", "Type", "PK", "FK →", "Nullable", "Default", "Unique",
                     "Thai label", "Source (page / JS property)", "Description"],
                    rows, [9, 28, 5, 28, 18, 5, 30, 9, 16, 7, 28, 45, 45], fills)

    e_rows = [[s["schema"], e["table"], r.get("code", ""), r.get("label_th", ""),
               "; ".join(f"{k}={v}" for k, v in (r.get("extra") or {}).items())]
              for s in schemas for e in s.get("enums", []) for r in e.get("rows", [])]
    write_sheet(wb.create_sheet("Enums"), ["Schema", "Lookup table", "Code", "Thai label", "Extra"],
                e_rows, [10, 28, 40, 55, 60])

    f_rows = [[s["activity"], i, f.get("page", ""), f.get("role", ""),
               ", ".join(f.get("reads", [])), ", ".join(f.get("writes", []))]
              for s in schemas for i, f in enumerate(s.get("flow_map", []), start=1)]
    write_sheet(wb.create_sheet("Flow map"), ["Activity", "Order", "Page", "Role", "Reads", "Writes"],
                f_rows, [9, 7, 48, 24, 55, 55])

    n_rows = [[s["activity"], i, n] for s in schemas for i, n in enumerate(s.get("notes", []), start=1)]
    write_sheet(wb.create_sheet("Notes"), ["Activity", "#", "Note"], n_rows, [9, 5, 140])
    wb.save(path)


# ---------------------------------------------------------------- sql
def sql_literal(text):
    return "'" + text.replace("'", "''") + "'"


def sql_value(v):
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return str(v)
    return sql_literal(str(v))


def enum_inserts(schema):
    """Seed rows for lookup tables; `extra` keys are kept only when they match a real column."""
    cols_of = {t["name"]: [c["name"] for c in t["columns"]] for t in schema["tables"]}
    out = []
    for e in schema.get("enums", []):
        cols = cols_of.get(e["table"])
        if not cols or not e.get("rows"):
            continue
        out.append(f"\n-- seed {schema['schema']}.{e['table']}")
        for r in e["rows"]:
            values = {"code": r.get("code"), "label_th": r.get("label_th")}
            values.update({k: v for k, v in (r.get("extra") or {}).items() if k in cols})
            names = [k for k in values if k in cols]
            out.append(f"insert into {schema['schema']}.{e['table']} ({', '.join(names)}) values "
                       f"({', '.join(sql_value(values[k]) for k in names)}) on conflict do nothing;")
    return out


def build_sql(schemas, path):
    out = ["-- Generated from schema-10-*.json — do not edit by hand."]
    for s in schemas:
        sch = s["schema"]
        out.append(f"\n-- ===== {s['activity']} {s.get('title_th', '')}\ncreate schema if not exists {sch};")
        fks = []
        # External tables (e.g. กิจกรรมที่ 7's tbl_board_submission) are documented, not created,
        # and references to them stay logical: no FK constraint across modules.
        external = {t["name"] for t in s["tables"] if t.get("external")}
        for t in s["tables"]:
            if t["name"] in external:
                out.append(f"\n-- {sch}: {t['name']} is owned by another module (reference only, not created here)")
                continue
            lines = []
            for c in t["columns"]:
                ctype = "bigint generated always as identity" if c["type"] == "bigserial" else c["type"]
                line = f"  {c['name']} {ctype}"
                if not c.get("nullable", True) and not c.get("pk"):
                    line += " not null"
                if c.get("default") is not None and c["type"] != "bigserial":
                    line += f" default {c['default']}"
                if c.get("unique"):
                    line += " unique"
                lines.append(line)
                if c.get("fk") and c["fk"].split(".")[0] not in external:
                    rt, _, rc = c["fk"].partition(".")
                    fks.append(f"alter table {sch}.{t['name']} add constraint fk_{t['name']}_{c['name']} "
                               f"foreign key ({c['name']}) references {sch}.{rt} ({rc or 'id'});")
            pks = [c["name"] for c in t["columns"] if c.get("pk")]
            if pks:
                lines.append(f"  primary key ({', '.join(pks)})")
            lines += [f"  unique ({', '.join(u)})" for u in t.get("unique_constraints", [])]
            lines += [f"  check ({ck})" for ck in t.get("check_constraints", [])]
            out.append(f"\ncreate table {sch}.{t['name']} (\n" + ",\n".join(lines) + "\n);")
            out.append(f"comment on table {sch}.{t['name']} is "
                       f"{sql_literal(t.get('label_th', '') + ' — ' + t.get('description', ''))};")
            out += [f"create index on {sch}.{t['name']} ({', '.join(ix)});" for ix in t.get("indexes", [])]
        out.append("")
        out += fks
        out += enum_inserts(s)
    path.write_text("\n".join(out) + "\n", encoding="utf-8")


# ---------------------------------------------------------------- checks
def validate(s):
    cols_of = {t["name"]: {c["name"] for c in t["columns"]} for t in s["tables"]}
    problems = []
    if len(cols_of) != len(s["tables"]):
        problems.append("duplicate table names")
    for t in s["tables"]:
        names = [c["name"] for c in t["columns"]]
        if len(names) != len(set(names)):
            problems.append(f"{t['name']}: duplicate column names")
        for c in t["columns"]:
            if not c.get("fk"):
                continue
            rt, _, rc = c["fk"].partition(".")
            if rt not in cols_of:
                problems.append(f"{t['name']}.{c['name']}: fk table '{rt}' missing")
            elif rc and rc not in cols_of[rt]:
                problems.append(f"{t['name']}.{c['name']}: fk column '{c['fk']}' missing")
    return problems


def main():
    schemas = [load(p) for p in INPUTS]
    for s in schemas:
        for problem in validate(s):
            print(f"  PROBLEM {s['schema']}: {problem}")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    edges = build_drawio(schemas, OUT_DIR / "ecmis-activity10-er.drawio")
    build_xlsx(schemas, OUT_DIR / "ecmis-activity10-data-dictionary.xlsx")
    build_sql(schemas, OUT_DIR / "ecmis-activity10-schema.sql")
    for s in schemas:
        print(f"{s['schema']}: {len(s['tables'])} tables, {sum(len(t['columns']) for t in s['tables'])} columns")
    print(f"fk edges: {edges}")


if __name__ == "__main__":
    main()
