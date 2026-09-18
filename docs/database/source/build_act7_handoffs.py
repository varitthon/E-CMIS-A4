"""Build the activity-7 send/receive summary workbook from act7-handoffs.json.

Usage: python build_act7_handoffs.py OUT.xlsx
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = Path(__file__).parent / "act7-handoffs.json"
OUT = Path(sys.argv[1])

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SEND_FILL = PatternFill("solid", fgColor="FCE4D6")
RECV_FILL = PatternFill("solid", fgColor="E2EFDA")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
LINK_FONT = Font(name=FONT, size=10, color="0563C1", underline="single")

CATEGORY_TH = {
    "PAGE_INPUT": "กรอกในหน้านี้",
    "PAGE_DECISION": "การตัดสินใจในหน้านี้",
    "RECEIVED_RESULT": "ผลที่ได้รับกลับ",
    "DOC_NUMBER": "เลขที่หนังสือ",
    "CARRIED_DOCUMENT": "เอกสารที่แนบไปด้วย",
    "CARRIED_SIGNATURE": "ลายมือชื่อที่แนบไปด้วย",
    "CARRIED_ATTACHMENT": "ไฟล์แนบที่แนบไปด้วย",
    "CASE_REF": "ข้อมูลอ้างอิงคดี",
}
# The 10.3 pages do label these handoffs กิจกรรมที่ 7 (10-3-09 and 10-3v-28 HTML); the
# extraction only searched ecmis-10-3.js, so its "no match" remarks are dropped.
STALE_ISSUE = re.compile(r"(no match|found no|not found|ไม่พบ)", re.I)


def head(ws, row, headers):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill, c.border = HEAD_FILL, BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")


def body(ws, row, values, fill=None):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.border, c.alignment = Font(name=FONT, size=10), BORDER, WRAP_TOP
        if fill:
            c.fill = fill


def set_widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def clean_issues(p):
    issues = p.get("issues", [])
    if p["activity"] != "10.3":
        return issues
    return [i for i in issues if not ("กิจกรรมที่ 7" in i and STALE_ISSUE.search(i))]


def is_this_step(ref):
    return ref["part"].startswith(("จุดนี้", "ส่งเข้า")) or "/ จุดนี้" in ref["part"]


def law_lines(refs):
    """One 'LAWxxxx ชื่อขั้นตอน' line per distinct code, in flow order."""
    seen = {}
    for x in refs:
        if x["code"] != "—":
            seen.setdefault(x["code"], x["text"])
    return "\n".join(f"{code} {text}" for code, text in seen.items())


def main_law(p):
    """LAW steps performed at this point itself (not the steps before/after it), with Thai labels."""
    return law_lines([x for x in p.get("law_refs", []) if is_this_step(x)])


def write_law_refs(ws, p, row):
    """TO-BE LAW reference table; returns the next free row."""
    ws.cell(row=row, column=1, value="อ้างอิงขั้นตอน LAW (TO-BE UserFlow V1.0.1)").font = Font(name=FONT, size=12, bold=True)
    head(ws, row + 1, ["รหัส LAW", "ขั้นตอน (ตาม TO-BE)", "", "ตำแหน่งเทียบกับจุดนี้", "หน้า drawio (TO-BE)"])
    ws.merge_cells(start_row=row + 1, start_column=2, end_row=row + 1, end_column=3)
    ws.merge_cells(start_row=row + 1, start_column=5, end_row=row + 1, end_column=10)
    r = row + 1
    for r, x in enumerate(p.get("law_refs", []), start=row + 2):
        is_here = x["part"].startswith(("จุดนี้", "ส่งเข้า")) or "/ จุดนี้" in x["part"]
        body(ws, r, [x["code"], x["text"], "", x["part"], x["tobe"], "", "", "", "", ""],
             PatternFill("solid", fgColor="DDEBF7") if is_here else None)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=10)
        if is_here:
            ws.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True)
    if p.get("law_note"):
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        c = ws.cell(row=r, column=1, value=p["law_note"])
        c.font, c.alignment = Font(name=FONT, size=10, italic=True), WRAP_TOP
        if p["law_note"].startswith("⚠"):
            c.fill = WARN_FILL
        ws.row_dimensions[r].height = 30
    return r + 2


def sheet_name(p, used):
    name = re.sub(r"[\[\]:*?/\\]", "-", p.get("sheet_title") or p["id"])[:31]
    base, n = name, 2
    while name in used:
        name = f"{base[:28]}-{n}"
        n += 1
    used.add(name)
    return name


def build_overview(wb, points, names):
    ws = wb.active
    ws.title = "Overview"
    ws["A1"] = "จุดส่งข้อมูลไป / รับข้อมูลจาก กิจกรรมที่ 7 — E-CMIS กิจกรรมที่ 10"
    ws["A1"].font = Font(name=FONT, size=16, bold=True)
    ws["A2"] = ("หนึ่งชีตต่อหนึ่งจุด (คลิก ID เพื่อเปิด) · สีส้ม = ส่งไปกิจกรรมที่ 7 · สีเขียว = รับกลับ · "
                "รหัส LAW อ้างอิงจาก 'กิจกรรมที่ 10 TO-BE 10.x-UserFlow V1.0.1.drawio' · "
                "สีเหลือง = มีหมายเหตุ/ช่องว่างในม็อกอัป หรือ TO-BE กับม็อกอัปไม่ตรงกัน")
    ws["A2"].font = Font(name=FONT, size=10, italic=True)
    ws["A2"].alignment = WRAP_TOP
    ws.merge_cells("A2:L2")
    ws.row_dimensions[2].height = 32
    head(ws, 4, ["ID", "กิจกรรม", "ทิศทาง", "LAW จุดนี้ (TO-BE)", "LAW ที่เกี่ยวข้องทั้งหมด", "ยืนยันว่าเป็นกิจกรรมที่ 7",
                 "การดำเนินการ", "หน้า (ม็อกอัป)", "ผู้ดำเนินการ", "สถานะก่อน", "สถานะหลัง", "คู่ส่ง/รับ",
                 "จำนวนรายการข้อมูล", "ประเด็นที่ต้องยืนยัน"])
    for r, p in enumerate(points, start=5):
        is_send = p["direction"] == "SEND"
        body(ws, r, [p["id"], p["activity"], "➡ ส่ง" if is_send else "⬅ รับ", main_law(p),
                     law_lines(p.get("law_refs", [])),
                     "ใช่" if p.get("confirmed") else "ยังไม่ยืนยัน", p.get("action_th", ""), p["page"],
                     p.get("role_th", p.get("role", "")), p.get("status_before", ""), p.get("status_after", ""),
                     p.get("pair_id", ""), len(p.get("data", [])), len(clean_issues(p)) + (1 if p.get("law_note", "").startswith("⚠") else 0)],
             SEND_FILL if is_send else RECV_FILL)
        link = ws.cell(row=r, column=1)
        link.hyperlink, link.font = f"#'{names[p['id']]}'!A1", LINK_FONT
        ws.cell(row=r, column=4).font = Font(name=FONT, size=10, bold=True)
        if p.get("law_note", "").startswith("⚠"):
            ws.cell(row=r, column=4).fill = WARN_FILL
        if not p.get("confirmed"):
            ws.cell(row=r, column=6).fill = WARN_FILL
    set_widths(ws, [11, 9, 9, 48, 58, 14, 45, 42, 26, 30, 30, 11, 11, 11])
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:N{4 + len(points)}"


def write_info(ws, p, names, fill):
    info = [
        ("การดำเนินการ", p.get("action_th", "")),
        ("หน้า", p["page"]),
        ("ผู้ดำเนินการ", f"{p.get('role_th', '')} ({p.get('role', '')})"),
        ("สถานะก่อน", p.get("status_before", "")),
        ("สถานะหลัง", p.get("status_after", "")),
        ("ขั้นตอนก่อนหน้า", p.get("previous_steps", "")),
        ("ขั้นตอนถัดไป", p.get("next_steps", "")),
        ("คู่ส่ง/รับ", p.get("pair_id", "")),
        ("LAW จุดนี้ (TO-BE)", main_law(p)),
        ("ยืนยันว่าเป็นกิจกรรมที่ 7",
         p.get("confirm_basis") or ("ใช่" if p.get("confirmed") else "ยังไม่ยืนยัน — โค้ดไม่ได้ระบุว่าเป็นกิจกรรมที่ 7")),
    ]
    for r, (key, value) in enumerate(info, start=3):
        kc = ws.cell(row=r, column=1, value=key)
        kc.font, kc.fill, kc.border, kc.alignment = Font(name=FONT, bold=True, size=10), fill, BORDER, WRAP_TOP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
        vc = ws.cell(row=r, column=2, value=value)
        vc.font, vc.border, vc.alignment = Font(name=FONT, size=10), BORDER, WRAP_TOP
        if isinstance(value, str) and "\n" in value:
            ws.row_dimensions[r].height = 16 * (value.count("\n") + 1)
        if key == "คู่ส่ง/รับ" and value in names:
            vc.hyperlink, vc.font = f"#'{names[value]}'!A1", LINK_FONT
        if key.startswith("ยืนยัน") and not p.get("confirmed"):
            vc.fill = WARN_FILL
    return 3 + len(info)


def build_point(wb, p, name, names):
    ws = wb.create_sheet(name)
    is_send = p["direction"] == "SEND"
    fill = SEND_FILL if is_send else RECV_FILL
    ws["A1"] = f"{p['id']} — {'ส่งไปกิจกรรมที่ 7' if is_send else 'รับจากกิจกรรมที่ 7'} (กิจกรรม {p['activity']})"
    ws["A1"].font = Font(name=FONT, size=15, bold=True)
    ws["J1"].value, ws["J1"].hyperlink, ws["J1"].font = "← Overview", "#'Overview'!A1", LINK_FONT

    start = write_law_refs(ws, p, write_info(ws, p, names, fill) + 1)
    ws.cell(row=start, column=1, value="ข้อมูลที่ส่ง" if is_send else "ข้อมูลที่รับ").font = Font(name=FONT, size=12, bold=True)
    head(ws, start + 1, ["#", "ประเภท", "รายการ", "Field id", "Property", "ชนิด", "บังคับ", "ตัวเลือก",
                         "ที่มา (หน้า/ขั้นตอน)", "หมายเหตุ"])
    order = list(CATEGORY_TH)
    rows = sorted(p.get("data", []), key=lambda d: order.index(d["category"]) if d["category"] in order else 99)
    for i, d in enumerate(rows, start=1):
        note = d.get("note", "")
        body(ws, start + 1 + i, [i, CATEGORY_TH.get(d["category"], d["category"]), d.get("label_th", ""),
                                 d.get("field_id", ""), d.get("property", ""), d.get("type", ""),
                                 d.get("required", ""), d.get("options", ""), d.get("origin", ""), note],
             WARN_FILL if note else None)

    issues = clean_issues(p)
    if issues:
        r = start + 3 + len(rows)
        ws.cell(row=r, column=1, value="ประเด็นที่ต้องยืนยัน / ข้อสังเกต").font = Font(name=FONT, size=12, bold=True)
        for j, text in enumerate(issues, start=1):
            ws.cell(row=r + j, column=1, value=j).font = Font(name=FONT, size=10)
            ws.merge_cells(start_row=r + j, start_column=2, end_row=r + j, end_column=10)
            c = ws.cell(row=r + j, column=2, value=text)
            c.font, c.alignment = Font(name=FONT, size=10), WRAP_TOP
            ws.row_dimensions[r + j].height = 15 * max(1, len(text) // 150 + 1)
    set_widths(ws, [22, 20, 34, 24, 28, 12, 8, 30, 30, 40])
    ws.freeze_panes = ws.cell(row=start + 2, column=1)


def main():
    points = json.loads(SRC.read_text(encoding="utf-8"))["points"]
    used, names = {"Overview"}, {}
    for p in points:
        names[p["id"]] = sheet_name(p, used)
    wb = Workbook()
    build_overview(wb, points, names)
    for p in points:
        build_point(wb, p, names[p["id"]], names)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"{len(points)} points, {sum(len(p.get('data', [])) for p in points)} data rows")


if __name__ == "__main__":
    main()
